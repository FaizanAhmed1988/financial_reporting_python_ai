"""
tests/test_acceptance.py
========================
UP6 — acceptance criteria across UP1..UP6, end to end against the real dataset.

Self-contained: builds its own malformed fixture, generates its own Excel export.
Run with:  python tests/test_acceptance.py
Exits non-zero if any criterion fails.
"""
import os, sys, json, subprocess
ROOT="/Volumes/ExternalSSD/AllProject/PersonalApp/financial_reporting_python_ai"
sys.path.insert(0, ROOT); os.chdir(ROOT)
import pandas as pd
from pathlib import Path
from streamlit.testing.v1 import AppTest
from src.data_loader import load_workbook, ConfigurableColumnMapper, DEFAULT_COLUMN_ALIASES
from src.account_classifier import AccountClassifier
from src.file_detector import detect_file_type
from src.validation_engine import validate_dataset
from src.coa_mapper import build_coa_comparison, load_reference_coa, load_user_mappings, make_source_key, known_account_codes
import src.dataset_manager as dm, src.insights_engine as ie, src.upload_history as uh
from src.security import sanitize_cell

R=[]
def check(cid, desc, ok, detail=""):
    R.append((cid, desc, ok, detail)); 
    print(f"  {'PASS' if ok else 'FAIL'}  {cid:<7} {desc}" + (f"  [{detail}]" if detail else ""))

WB="data/raw/GL_COA_TB_Dummy_Dataset.xlsx"
wb=load_workbook(WB,header_row=3); coa=AccountClassifier().classify(wb.coa)
orig=dm.build_original_dataset(coa,wb); S=dm.compute_statements(orig)

print("\n— BASELINE INTEGRITY —")
check("BASE-1","Unit tests 6/6",
      subprocess.run([f"{ROOT}/.venv/bin/python","tests/test_financials.py"],capture_output=True).returncode==0)
check("BASE-2","Revenue 897,000", abs(S["pl_metrics"]["Total Revenue"]-897000)<.01)
check("BASE-3","Net loss (1,216,800)", abs(S["pl_metrics"]["Net Profit"]+1216800)<.01)
check("BASE-4","Total assets 10,509,300", abs(S["bs_metrics"]["Total Assets"]-10509300)<.01)
check("BASE-5","Closing cash 960,100", abs(S["cf_metrics"]["Closing Cash (Calculated)"]-960100)<.01)
check("BASE-6","Balance sheet balances", S["bs_metrics"]["Is Balanced"])
check("BASE-7","Cash flow reconciles", S["cf_metrics"]["Reconciled"])
check("BASE-8","0 unclassified accounts", int(coa["financial_statement"].isna().sum())==0)

print("\n— UP1 detection & mapping —")
gl_raw=pd.read_excel(WB,sheet_name="General Ledger",header=3,dtype=str).dropna(how="all")
gl_raw=gl_raw[gl_raw["Account Code"].notna()]
d,c,rm=detect_file_type(gl_raw,ConfigurableColumnMapper(),"Auto Detect")
check("UP1-1","GL detected at 100%", d=="General Ledger" and c==1.0, f"{d} {c*100:.0f}%")
tb_raw=pd.read_excel(WB,sheet_name="Trial Balance",header=3,dtype=str).dropna(how="all")
tb_raw=tb_raw[tb_raw["Account Code"].notna()]
d2,c2,rm2=detect_file_type(tb_raw,ConfigurableColumnMapper(),"Auto Detect")
check("UP1-2","TB detected at 100%", d2=="Trial Balance" and c2==1.0, f"{d2} {c2*100:.0f}%")
m=ConfigurableColumnMapper()
check("UP1-3","camelCase headers map (H3)", m.find_canonical("AccountNumber")=="account_code")
bad=[v for cn,vs in DEFAULT_COLUMN_ALIASES.items() for v in vs if m.find_canonical(v)!=cn and v.lower()!="description"]
check("UP1-4","86/86 aliases unchanged", not bad, f"{len(bad)} broken")

print("\n— UP2 validation & COA mapping —")
ref,avail,msg=load_reference_coa(None); store=load_user_mappings()
glm=gl_raw.rename(columns=rm); sk=make_source_key("gl","general_ledger")
rep=validate_dataset(glm,"general_ledger",known_account_codes(ref,store,sk),avail,[])
check("UP2-1","Clean GL validates PASS", rep.status=="PASS", rep.status)
messy=glm.head(12).copy().astype(object)
messy.loc[messy.index[1],"credit"]="900.0"; messy.loc[messy.index[3],"account_code"]="9999"
messy.loc[messy.index[5],"date"]="not-a-date"; messy.loc[messy.index[6],"debit"]="TWELVE THOUSAND"
messy.loc[messy.index[7],"account_code"]=""; messy.loc[messy.index[9],"credit"]="0"
messy.loc[messy.index[9],"debit"]="0"
messy=pd.concat([messy,messy.tail(1)],ignore_index=True)
mrep=validate_dataset(messy,"general_ledger",known_account_codes(ref,store,sk),avail,["9999"])
check("UP2-2","Messy GL blocked with 7 errors", mrep.status=="ERROR" and len(mrep.errors)==7, f"{len(mrep.errors)} errors")
cmp_=build_coa_comparison(glm,"general_ledger",ref,store,sk,avail,msg)
check("UP2-3","26/26 GL accounts matched", cmp_.summary["Matched"]==26, str(cmp_.summary["Matched"]))
check("UP2-4","Never auto-assigns unknown accounts",
      all(r["Assigned Category"]=="— Mapping Required —" for _,r in cmp_.table.iterrows() if r["Status"]=="Unmapped"))

print("\n— UP3 dataset manager —")
ds_tb=dm.build_separate_dataset("TB","trial_balance",tb_raw.rename(columns=rm2),coa,file_name="tb.xlsx")
s_tb=dm.compute_statements(ds_tb)
check("UP3-1","TB upload reproduces original exactly",
      abs(s_tb["pl_metrics"]["Net Profit"]+1216800)<.01 and abs(s_tb["bs_metrics"]["Total Assets"]-10509300)<.01)
dup=dm.detect_duplicates(orig.gl,glm)
check("UP3-2","Re-upload = 200 confirmed duplicates, 0 new",
      dup.counts["Confirmed Duplicates"]==200 and dup.counts["New Transactions"]==0, str(dup.counts))
q2=glm.head(6).copy(); q2["date"]="2026-04-15"; q2["voucher_no"]=["JV-9001"]*2+["JV-9002"]*2+["JV-9003"]*2
mixed=pd.concat([q2,glm.head(3)],ignore_index=True); d2_=dm.detect_duplicates(orig.gl,mixed)
check("UP3-3","Mixed file = 6 new / 3 confirmed",
      d2_.counts["New Transactions"]==6 and d2_.counts["Confirmed Duplicates"]==3, str(d2_.counts))
app=dm.build_appended_dataset("Combined",orig,d2_.new_rows)
check("UP3-4","Append rebuilds a balanced TB",
      abs(app.tb["closing_debit"].sum()-app.tb["closing_credit"].sum())<.01)
after=dm.compute_statements(orig)
check("UP3-5","ORIGINAL IMMUTABLE after all operations",
      abs(after["pl_metrics"]["Net Profit"]+1216800)<.01 and len(orig.gl)==200)
check("UP3-6","Append refused for TB slot","trial_balance" not in dm.APPENDABLE_SLOTS)

print("\n— UP4 history & duplicate prevention —")
h=uh._empty_store(); fp=uh.fingerprint(b"payload")
h=uh.record(h,uh.new_entry(file_name="a.xlsx",file_hash=fp,outcome=uh.OUTCOME_PROCESSED,dataset_name="D1"))
check("UP4-1","Renamed copy detected by content hash", uh.find_prior_uploads(h,uh.fingerprint(b"payload")).was_processed)
check("UP4-2","Different sheet is NOT a duplicate", not uh.find_prior_uploads(h,fp,sheet="Other").seen)
check("UP4-3","Audit trail renders", not uh.to_dataframe(h).empty)

print("\n— UP5 insights & error handling —")
ins=ie.generate_insights(S,orig.limitations,0,orig.name)
check("UP5-1","Insights generated with evidence", len(ins)>=10 and all(i.explanation for i in ins), f"{len(ins)} findings")
check("UP5-2","Salary driver computed correctly (117%, not 170%)",
      any("117% of revenue" in i.explanation for i in ins))
gl_only=dm.build_separate_dataset("GLonly","general_ledger",glm,coa)
i2=ie.generate_insights(dm.compute_statements(gl_only),gl_only.limitations,0,"GLonly")
check("UP5-3","Position insights suppressed without opening balances",
      not any(x.category in (ie.CAT_LIQUIDITY,ie.CAT_WORKING_CAPITAL,ie.CAT_CASH) for x in i2))
check("UP5-4","Profitability insights still fire on GL-only",
      any(x.category==ie.CAT_PROFITABILITY for x in i2))

print("\n— UP6 export & security —")
check("UP6-1","Formula payload neutralised", str(sanitize_cell("=cmd|'/c calc'!A1")).startswith("'"))
check("UP6-2","Negative numbers NOT corrupted", sanitize_cell("-1,234.50")=="-1,234.50")
import tempfile
from src.ar_analysis import analyze_ar
from src.ap_analysis import analyze_ap
from src.budget_vs_actual import generate_bva
from src.controls import run_control_tests
from src.anomaly_detection import run_ml_anomaly_detection
from src.forecasting import generate_forecast
from src.scenario_analysis import run_scenarios
from src.reporting import generate_excel_report
ex=Path(tempfile.mkdtemp())/"acceptance_export.xlsx"
generate_excel_report(orig.coa_classified,orig.gl,orig.tb,S["pl_df"],S["pl_metrics"],
    S["bs_df"],S["bs_metrics"],S["cf_df"],S["cf_metrics"],S["ratios_df"],S["wc_df"],
    analyze_ar(S["bs_df"],orig.tb,orig.coa_classified),analyze_ap(S["bs_df"],orig.tb,orig.coa_classified),
    generate_bva(S["pl_df"],pd.read_csv("data/synthetic/budget_q1_2026.csv")),
    run_control_tests(orig.gl),run_ml_anomaly_detection(orig.gl),
    generate_forecast(orig.tb,orig.coa_classified,orig.gl),
    run_scenarios(S["pl_metrics"],S["bs_metrics"],S["bs_df"]),
    output_path=ex,dataset_name=orig.name,limitations=orig.limitations)
if ex.exists():
    from openpyxl import load_workbook as xl
    w=xl(ex); check("UP6-3","Export has 15 sheets", len(w.sheetnames)==15, str(len(w.sheetnames)))
    rows=[str(w["15. Management Summary"].cell(row=r,column=1).value or "") for r in range(1,60)]
    check("UP6-4","Summary titled by dataset", any("Original Dataset (Q1 2026)" in r for r in rows))
    check("UP6-5","Hardcoded 170% claim removed", not any("170%" in r for r in rows))
for f,ok in [("data/raw/GL_COA_TB_Dummy_Dataset.xlsx",1),("data/processed/upload_history.json",1)]:
    g=subprocess.run(["git","check-ignore","-q",f]).returncode==0
    check(f"UP6-6","Financial data gitignored: "+Path(f).name, g)
net=subprocess.run(["grep","-rEl","requests\\.|urllib|openai|anthropic|api_key","src/"],capture_output=True,text=True).stdout.strip()
check("UP6-7","No outbound network calls / API keys in src/", not net, net or "clean")

print("\n— DASHBOARD —")
at=AppTest.from_file(f"{ROOT}/dashboard/app.py",default_timeout=240); at.run()
check("DASH-1","App boots without exception", not at.exception)
secs=list(at.sidebar.radio[0].options)
check("DASH-2","All 8 sections present and unchanged",
      secs==["Executive Summary","Financial Statements","Ratios & Working Capital","Sub-Ledgers",
             "Controls & AI","Scenario Analysis","Month-End Close","Upload Financial Data"])
bad=[]
for n in secs:
    at.sidebar.radio[0].set_value(n).run()
    if at.exception or at.error: bad.append(n)
check("DASH-3","All 8 sections render cleanly", not bad, str(bad))
at.sidebar.radio[0].set_value("Executive Summary").run()
mets={m.label:m.value for m in at.metric}
check("DASH-4","Original figures correct on screen",
      mets.get("Total Revenue")=="897,000" and mets.get("Net Profit / (Loss)")=="-1,216,800")
check("DASH-5","Insights panel live", "🔴 Critical" in mets, mets.get("🔴 Critical",""))

p_=sum(1 for _,_,o,_ in R if o); t=len(R)
print(f"\n{'='*70}\nACCEPTANCE: {p_}/{t} passed")
if p_<t:
    print("FAILURES:"); [print(f"  {c} — {d} {x}") for c,d,o,x in R if not o]
print("="*70)
sys.exit(0 if p_==t else 1)
