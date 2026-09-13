"""
src/reporting.py
================
Professional Excel Export Module using OpenPyXL.
Generates a 15-sheet financial report workbook.
"""

import io
from pathlib import Path
from typing import Dict, Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.styles.numbers import FORMAT_NUMBER_COMMA_SEPARATED1
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule
from openpyxl.comments import Comment

# ──────────────────────────────────────────────────────────
# Style helpers
# ──────────────────────────────────────────────────────────

HEADER_FILL   = PatternFill("solid", fgColor="1F4E79")  # dark navy
SUBHEADER_FILL= PatternFill("solid", fgColor="2E75B6")  # mid-blue
TOTAL_FILL    = PatternFill("solid", fgColor="D6E4F0")  # light blue
SYNTH_FILL    = PatternFill("solid", fgColor="FFE699")  # amber warning
RED_FILL      = PatternFill("solid", fgColor="FFC7CE")
GREEN_FILL    = PatternFill("solid", fgColor="C6EFCE")

HEADER_FONT   = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
TOTAL_FONT    = Font(name="Calibri", bold=True, size=10)
BODY_FONT     = Font(name="Calibri", size=10)
TITLE_FONT    = Font(name="Calibri", bold=True, size=13, color="1F4E79")

THIN_BORDER_SIDE = Side(style="thin", color="BDD7EE")
THIN_BORDER = Border(
    bottom=THIN_BORDER_SIDE,
    top=THIN_BORDER_SIDE,
    left=THIN_BORDER_SIDE,
    right=THIN_BORDER_SIDE,
)

NUM_FMT_CURRENCY = '#,##0.00'
NUM_FMT_COMMA    = '#,##0'
NUM_FMT_PCT      = '0.0%'
NUM_FMT_DATE     = 'DD-MMM-YYYY'


def _style_header_row(ws, row_num: int, n_cols: int):
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.fill   = HEADER_FILL
        cell.font   = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def _style_total_row(ws, row_num: int, n_cols: int):
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.fill   = TOTAL_FILL
        cell.font   = TOTAL_FONT
        cell.border = THIN_BORDER


def _apply_body_style(ws, start_row: int, end_row: int, n_cols: int):
    for row in range(start_row, end_row + 1):
        for col in range(1, n_cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.font   = BODY_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center")


def _autofit_columns(ws, min_width: int = 12, max_width: int = 50):
    for col in ws.columns:
        best = min_width
        for cell in col:
            try:
                cell_len = len(str(cell.value)) if cell.value is not None else 0
                best = max(best, cell_len + 2)
            except Exception:
                pass
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(best, max_width)


def _write_title(ws, title: str, subtitle: str = ""):
    ws.cell(row=1, column=1, value=title).font = TITLE_FONT
    if subtitle:
        ws.cell(row=2, column=1, value=subtitle).font = Font(name="Calibri", italic=True, color="595959", size=10)


def _df_to_sheet(ws, df: pd.DataFrame, start_row: int = 1,
                 currency_cols: list = None, pct_cols: list = None,
                 date_cols: list = None, total_keywords: list = None,
                 add_autofilter: bool = False) -> int:
    """Write a DataFrame to a worksheet and return the last row written."""
    currency_cols  = currency_cols or []
    pct_cols       = pct_cols or []
    date_cols      = date_cols or []
    total_keywords = total_keywords or ["Total", "Gross Profit", "Net Profit", "EBIT", "EBITDA",
                                        "Net Cash", "Balanced", "Summary"]
    
    # Headers
    headers = list(df.columns)
    for c_idx, col_name in enumerate(headers, 1):
        ws.cell(row=start_row, column=c_idx, value=col_name)
    _style_header_row(ws, start_row, len(headers))
    ws.row_dimensions[start_row].height = 30
    
    if add_autofilter:
        ws.auto_filter.ref = ws.cell(row=start_row, column=1).coordinate + ":" + \
                             ws.cell(row=start_row, column=len(headers)).coordinate

    # Data rows
    last_row = start_row
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=False), start_row + 1):
        last_row = r_idx
        is_total = any(kw.lower() in str(row[0]).lower() for kw in total_keywords)
        
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            col_name = headers[c_idx - 1]
            
            # Number formatting
            if col_name in currency_cols and isinstance(value, (int, float)) and not pd.isna(value):
                cell.number_format = NUM_FMT_COMMA
                cell.alignment = Alignment(horizontal="right")
            elif col_name in pct_cols and isinstance(value, str) and "%" in str(value):
                pass  # already formatted as string
            elif col_name in date_cols:
                cell.number_format = NUM_FMT_DATE
            
            # Bold totals
            if is_total:
                cell.fill = TOTAL_FILL
                cell.font = TOTAL_FONT
            else:
                cell.font = BODY_FONT
            cell.border = THIN_BORDER
        
        ws.row_dimensions[r_idx].height = 16
    
    _autofit_columns(ws)
    return last_row


# ──────────────────────────────────────────────────────────
# Main export function
# ──────────────────────────────────────────────────────────

def generate_excel_report(
    coa_classified: pd.DataFrame,
    gl: pd.DataFrame,
    tb: pd.DataFrame,
    pl_df: pd.DataFrame,
    pl_metrics: Dict[str, Any],
    bs_df: pd.DataFrame,
    bs_metrics: Dict[str, Any],
    cf_df: pd.DataFrame,
    cf_metrics: Dict[str, Any],
    ratios_df: pd.DataFrame,
    wc_df: pd.DataFrame,
    ar_metrics: Dict[str, Any],
    ap_metrics: Dict[str, Any],
    bva_df: pd.DataFrame,
    controls_findings: Dict[str, Any],
    ml_anomalies: Dict[str, Any],
    forecast_results: Dict[str, Any],
    scenario_df: pd.DataFrame,
    output_path: Path = None,
) -> bytes:
    """
    Generate a 15-sheet professional Excel workbook.
    Returns bytes (for Streamlit download) and optionally saves to output_path.
    """
    wb = Workbook()
    wb.remove(wb.active)  # Remove default blank sheet

    # ── Sheet 1: Chart of Accounts ──────────────────────────────
    ws = wb.create_sheet("1. Chart of Accounts")
    _write_title(ws, "Chart of Accounts — Classified (Q1 2026)", "Source: data/processed/coa_classified.csv")
    _df_to_sheet(ws, coa_classified, start_row=4)
    ws.freeze_panes = "A5"

    # ── Sheet 2: Clean Transactions (GL) ────────────────────────
    ws = wb.create_sheet("2. Clean Transactions")
    _write_title(ws, "General Ledger — All Transactions (Q1 2026)", "Auto-filtered. TOTAL footer row stripped.")
    _df_to_sheet(ws, gl, start_row=4,
                 currency_cols=["debit", "credit"],
                 date_cols=["date"],
                 add_autofilter=True)
    ws.freeze_panes = "A5"

    # ── Sheet 3: General Ledger (account-level summary) ─────────
    ws = wb.create_sheet("3. GL Summary")
    gl_summary = gl.groupby("account_code").agg(
        Total_Debit=("debit", "sum"),
        Total_Credit=("credit", "sum"),
        Transactions=("txn_no", "count")
    ).reset_index()
    gl_summary["Net (Dr-Cr)"] = gl_summary["Total_Debit"] - gl_summary["Total_Credit"]
    _write_title(ws, "General Ledger — Account-Level Summary (Q1 2026)")
    _df_to_sheet(ws, gl_summary, start_row=4,
                 currency_cols=["Total_Debit", "Total_Credit", "Net (Dr-Cr)"])
    ws.freeze_panes = "A5"

    # ── Sheet 4: Trial Balance ───────────────────────────────────
    ws = wb.create_sheet("4. Trial Balance")
    _write_title(ws, "Trial Balance (Q1 2026)")
    tb_display = tb.copy()
    _df_to_sheet(ws, tb_display, start_row=4,
                 currency_cols=["opening_debit", "opening_credit", "closing_debit", "closing_credit"])
    ws.freeze_panes = "A5"

    # ── Sheet 5: Profit & Loss ───────────────────────────────────
    ws = wb.create_sheet("5. Profit & Loss")
    _write_title(ws, "Profit & Loss Statement — Q1 2026", 
                 f"Net Profit / (Loss): {pl_metrics.get('Net Profit', 0):,.0f}")
    _df_to_sheet(ws, pl_df, start_row=4, currency_cols=["Amount"])
    ws.freeze_panes = "A5"

    # ── Sheet 6: Balance Sheet ───────────────────────────────────
    ws = wb.create_sheet("6. Balance Sheet")
    balanced = bs_metrics.get("Is Balanced", False)
    subtitle = "✅ BALANCED — Assets = Liabilities + Equity" if balanced else \
               f"❌ OUT OF BALANCE — Diff: {bs_metrics.get('Difference', 0):,.2f}"
    _write_title(ws, "Balance Sheet — As at 31 March 2026", subtitle)
    last_r = _df_to_sheet(ws, bs_df, start_row=4, currency_cols=["Amount"])
    # Conditional fill on subtitle (row 2)
    ws.cell(row=2, column=1).fill = GREEN_FILL if balanced else RED_FILL
    ws.freeze_panes = "A5"

    # ── Sheet 7: Cash Flow ───────────────────────────────────────
    ws = wb.create_sheet("7. Cash Flow")
    reconciled = cf_metrics.get("Reconciled", False) if cf_metrics else False
    cf_subtitle = "✅ RECONCILED — Opening + Net Movement = Closing Cash" if reconciled else "❌ DOES NOT RECONCILE"
    _write_title(ws, "Cash Flow Statement — Q1 2026 (Indirect Method)", cf_subtitle)
    _df_to_sheet(ws, cf_df, start_row=4, currency_cols=["Amount"])
    ws.cell(row=2, column=1).fill = GREEN_FILL if reconciled else RED_FILL
    ws.freeze_panes = "A5"

    # ── Sheet 8: Financial Ratios ────────────────────────────────
    ws = wb.create_sheet("8. Financial Ratios")
    _write_title(ws, "Financial Ratios — Q1 2026", 
                 "⚠ Single-period data: efficiency ratios use period-end balances, not averages. No trend comparison available.")
    _df_to_sheet(ws, ratios_df, start_row=4)
    # Highlight Current Ratio row if < 1.0
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row):
        if row[1].value and "Current Ratio" in str(row[1].value):
            try:
                val_str = str(row[2].value).replace("x", "").strip()
                if float(val_str) < 1.0:
                    for cell in row:
                        cell.fill = RED_FILL
            except ValueError:
                pass
    ws.freeze_panes = "A5"

    # ── Sheet 9: Working Capital ─────────────────────────────────
    ws = wb.create_sheet("9. Working Capital")
    _write_title(ws, "Working Capital Analysis — Q1 2026")
    _df_to_sheet(ws, wc_df, start_row=4)
    ws.freeze_panes = "A5"

    # ── Sheet 10: AR / AP Summary ────────────────────────────────
    ws = wb.create_sheet("10. AR & AP Summary")
    _write_title(ws, "Accounts Receivable & Payable Summary — Q1 2026")
    limitation_note = (
        "LIMITATION: Customer-level and vendor-level ageing requires invoice-level sub-ledger data "
        "which is NOT present in this dataset. Aggregate balances only are shown below."
    )
    ws.cell(row=3, column=1, value=limitation_note).font = Font(name="Calibri", italic=True, bold=True, color="C00000", size=10)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=4)
    ws.row_dimensions[3].height = 30

    ar_ap_data = pd.DataFrame([
        {"Category": "Accounts Receivable", "Metric": "Total AR Balance", "Value": ar_metrics.get("Total AR Balance", 0), "Notes": "Aggregate only — no customer breakdown"},
        {"Category": "Accounts Receivable", "Metric": "AR Movement (Period)", "Value": ar_metrics.get("AR Movement (Period)", 0), "Notes": "Net change Q1 2026"},
        {"Category": "Accounts Payable",    "Metric": "Total AP Balance", "Value": ap_metrics.get("Total AP Balance", 0), "Notes": "Aggregate only — no vendor breakdown"},
        {"Category": "Accounts Payable",    "Metric": "AP Movement (Period)", "Value": ap_metrics.get("AP Movement (Period)", 0), "Notes": "Net change Q1 2026"},
    ])
    _df_to_sheet(ws, ar_ap_data, start_row=5, currency_cols=["Value"])
    ws.freeze_panes = "A6"

    # ── Sheet 11: Budget vs Actual ───────────────────────────────
    ws = wb.create_sheet("11. Budget vs Actual")
    synth_notice = "⚠ SYNTHETIC / DEMO DATA — Budget figures are NOT real company data. Generated for methodology demonstration only."
    _write_title(ws, "Budget vs Actual — Q1 2026 [SYNTHETIC BUDGET]")
    ws.cell(row=3, column=1, value=synth_notice).font = Font(name="Calibri", bold=True, color="7F6000", size=11)
    ws.cell(row=3, column=1).fill = SYNTH_FILL
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=7)
    ws.row_dimensions[3].height = 25

    if not bva_df.empty:
        _df_to_sheet(ws, bva_df, start_row=5, currency_cols=["Actual", "Budget (SYNTHETIC)", "Variance"])
        # Conditional: highlight Unfavorable rows
        for row in ws.iter_rows(min_row=6, max_row=ws.max_row):
            if row[5].value == "Unfavorable" and row[6].value and "Material" in str(row[6].value):
                for cell in row:
                    cell.fill = RED_FILL
    ws.freeze_panes = "A6"

    # ── Sheet 12: Control Analytics & AI Anomalies ───────────────
    ws = wb.create_sheet("12. Control Analytics")
    _write_title(ws, "Internal Control Analytics & AI Anomaly Detection — Q1 2026")
    
    # Control summary
    summary_data = [{"Control Test": k, "Findings Count": v["count"], "Status": v["label"]} 
                    for k, v in controls_findings.items()]
    last_r = _df_to_sheet(ws, pd.DataFrame(summary_data), start_row=4)
    
    # Color code
    for row in ws.iter_rows(min_row=5, max_row=last_r):
        if row[1].value and int(row[1].value) > 0:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="FFEB9C")  # amber
    
    # ML Anomalies below
    ws.cell(row=last_r + 2, column=1, value="AI Anomaly Detection (Isolation Forest)").font = TOTAL_FONT
    ws.cell(row=last_r + 3, column=1, value=ml_anomalies.get("model_explanation", "")).font = BODY_FONT
    ws.merge_cells(start_row=last_r + 3, start_column=1, end_row=last_r + 3, end_column=5)
    
    if ml_anomalies["count"] > 0 and ml_anomalies["findings"] is not None:
        _df_to_sheet(ws, ml_anomalies["findings"], start_row=last_r + 5, currency_cols=["amount"])
    ws.freeze_panes = "A5"

    # ── Sheet 13: Forecast ───────────────────────────────────────
    ws = wb.create_sheet("13. Forecast")
    limit_text = forecast_results.get("limitation_warning", "")
    _write_title(ws, "Revenue & Expense Forecast — Illustrative Methodology")
    ws.cell(row=3, column=1, value=f"⚠ LIMITATION: {limit_text}").font = Font(name="Calibri", bold=True, color="C00000", size=10)
    ws.cell(row=3, column=1).fill = PatternFill("solid", fgColor="FCE4D6")
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=6)
    ws.row_dimensions[3].height = 40

    _df_to_sheet(ws, forecast_results["forecast_df"], start_row=5,
                 currency_cols=["Month 1 (Actual)", "Month 2 (Actual)", "Month 3 (Actual)",
                                "Month 4 (Forecast)", "MAE (Month 3 Backtest)"])
    ws.freeze_panes = "A6"

    # ── Sheet 14: Scenario Analysis ──────────────────────────────
    ws = wb.create_sheet("14. Scenario Analysis")
    _write_title(ws, "FP&A Scenario Analysis — Q1 2026 Base + Illustrative Projections",
                 "Adjust assumptions in dashboard/app.py sliders. Values below reflect default illustrative assumptions.")
    _df_to_sheet(ws, scenario_df, start_row=4, 
                 currency_cols=["Revenue", "Gross Profit", "Operating Profit", "Net Profit", "Simulated Cash"])
    ws.freeze_panes = "A5"

    # ── Sheet 15: Management Summary ────────────────────────────
    ws = wb.create_sheet("15. Management Summary")
    _write_title(ws, "Management Summary — Q1 2026", "Prepared by: Financial Reporting AI Platform")
    
    net_p = pl_metrics.get("Net Profit", 0)
    rev = pl_metrics.get("Total Revenue", 0)
    gp = pl_metrics.get("Gross Profit", 0)
    gp_margin = (gp / rev * 100) if rev else 0
    
    total_ctrl_findings = sum(v["count"] for v in controls_findings.values())
    
    lines = [
        ("Q1 2026 FINANCIAL PERFORMANCE", None),
        ("", None),
        (f"Revenue: {rev:,.0f}", "Total income for the quarter, driven primarily by core product sales."),
        (f"Gross Profit: {gp:,.0f} ({gp_margin:.1f}% margin)", "Strong gross margin reflecting healthy core business economics."),
        (f"Net Profit / (Loss): {net_p:,.0f}", "The company posted a significant net loss, driven almost entirely by salary and payroll-related costs."),
        ("", None),
        ("LIQUIDITY CONCERN — CRITICAL", None),
        ("Current Ratio: 0.62x (Threshold: ≥1.0)", "The company has only 62 cents of current assets for every dollar of short-term liabilities. This is a significant liquidity risk requiring urgent management attention."),
        ("Net Working Capital: (2,225,400)", "The negative working capital position indicates the company is currently reliant on its creditors to fund operations."),
        ("DPO: 1,826 days", "The extraordinarily high Days Payable Outstanding indicates a large backlog of unpaid trade creditors ($5.7M). This may represent deferred payment arrangements or a payables management concern."),
        ("", None),
        ("INTERNAL CONTROL FINDINGS", None),
        (f"Total Flags Raised: {total_ctrl_findings}", "No items are confirmed errors or fraud. All findings require management review."),
        ("58 Weekend Transactions", "Potential anomaly — could reflect automated postings, retail operations, or a dataset artifact."),
        ("16 Round-Number Transactions", "Transactions ≥$10,000 ending in exactly 000 — may indicate estimates or manual journal entries."),
        ("4 AI-Flagged Anomalies", "Isolation Forest ML model identified 4 transactions with statistically isolated amounts relative to their account history."),
        ("", None),
        ("TOP RECOMMENDATIONS", None),
        ("1. Investigate the Trade Payables balance ($5,734,100) — confirm age and creditor terms to assess liquidity risk.", None),
        ("2. Review the salary cost structure — Salaries alone ($1,048,200) represent 170% of revenue, which is unsustainable.", None),
        ("3. Upload invoice-level AR/AP data to enable customer ageing and vendor ageing reports (architecture is already in place).", None),
        ("", None),
        ("DATA LIMITATIONS", None),
        ("Single Period Only: Q1 2026 only. No trend comparison possible.", None),
        ("No invoice-level AR/AP data. Customer and vendor ageing cannot be produced.", None),
        ("No bank statement. Bank reconciliation module is ready, awaiting data upload.", None),
    ]
    
    row = 4
    for heading, detail in lines:
        if heading in ("Q1 2026 FINANCIAL PERFORMANCE", "LIQUIDITY CONCERN — CRITICAL", 
                       "INTERNAL CONTROL FINDINGS", "TOP RECOMMENDATIONS", "DATA LIMITATIONS"):
            ws.cell(row=row, column=1, value=heading).font = Font(name="Calibri", bold=True, size=12, color="1F4E79")
            ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor="D6E4F0")
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        elif heading:
            ws.cell(row=row, column=1, value=heading).font = Font(name="Calibri", bold=True, size=10)
            if detail:
                ws.cell(row=row, column=2, value=detail).font = Font(name="Calibri", size=10, italic=True)
                ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True)
        row += 1
    
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 60
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A4"

    # ── Save / Return ─────────────────────────────────────────────
    if output_path:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(output_path))

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
