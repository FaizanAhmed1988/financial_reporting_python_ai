import openpyxl
import pandas as pd

wb_path = r"data/raw/GL_COA_TB_Dummy_Dataset.xlsx"
wb = openpyxl.load_workbook(wb_path, read_only=True, data_only=True)
print("Sheets:", wb.sheetnames)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print("")
    print("=== SHEET:", sheet_name, "===")
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        vals = []
        for c in row:
            if c is None:
                vals.append("NULL")
            else:
                s = str(c)
                vals.append(s[:45])
        print("  Row", i+1, ":", vals)
        if i >= 9:
            print("  ...(showing first 10 rows only)")
            break

wb.close()
